"""Read-only OOXML extraction. Only a small arithmetic AST is interpreted, never eval/Excel code."""
import ast, datetime, hashlib, json, re
from decimal import Decimal
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'source/izvor-2017.xlsx'
w=openpyxl.load_workbook(p,data_only=False); cached=openpyxl.load_workbook(p,data_only=True)
now=datetime.datetime.now(datetime.timezone.utc).isoformat()
def val(v):
 if isinstance(v,(datetime.datetime,datetime.date)): return v.isoformat()
 return v
def arithmetic(x):
 def walk(n):
  if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)): return Decimal(str(n.value))
  if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)): return walk(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
  if isinstance(n,ast.BinOp):
   a,b=walk(n.left),walk(n.right)
   if isinstance(n.op,ast.Add):return a+b
   if isinstance(n.op,ast.Sub):return a-b
   if isinstance(n.op,ast.Mult):return a*b
   if isinstance(n.op,ast.Div):return a/b
  raise ValueError('Unsupported arithmetic')
 return walk(ast.parse(x,mode='eval').body)
def calc(sheet,coord,override=None,seen=None):
 coord=coord.replace('$',''); override=override or {}; seen=seen or set()
 if coord in override:return Decimal(str(override[coord]))
 if coord in seen:raise ValueError('Cycle')
 x=w[sheet][coord].value
 if x is None:return Decimal(0) # Historical regression only, NOT approval of a missing price.
 if isinstance(x,(int,float)):return Decimal(str(x))
 if not isinstance(x,str) or not x.startswith('='):raise ValueError('Not numeric')
 x=x[1:].replace('$','')
 if '[' in x or '!' in x:raise ValueError('External/cross-sheet')
 def summed(m):
  a,b=m.group(1),m.group(2); cells=w[sheet][a:b]
  return str(sum((calc(sheet,c.coordinate,override,seen|{coord}) for row in cells for c in row),Decimal(0)))
 x=re.sub(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)',summed,x,flags=re.I)
 x=re.sub(r'\b[A-Z]+\d+\b',lambda m:str(calc(sheet,m[0],override,seen|{coord})),x)
 return arithmetic(x)
def number(sheet,c,override=None):
 try:return str(calc(sheet,c,override)) if w[sheet][c].value is not None else None
 except (ValueError,ZeroDivisionError,ArithmeticError):return None
inventory=[]; issues=[]; cards=[]
for s in w:
 cells=[]
 for row in s:
  for c in row:
   if c.value is None and not c.comment:continue
   formula=c.value if c.data_type=='f' else None
   entry={'cell':c.coordinate,'original':val(c.value),'cached':val(cached[s.title][c.coordinate].value),'formula':formula,'dependencies':re.findall(r"(?:'[^']+'!|[A-Za-zČĆŽŠĐčćžšđ]+!)?\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?",formula or ''),'comment':c.comment.text if c.comment else None}
   cells.append(entry)
   if formula and '[' in formula:issues.append({'sheet':s.title,'cell':c.coordinate,'kind':'external','detail':formula,'resolution':'Povezati sa ostvarenom površinom/prinosom gazdinstva.'})
   if c.data_type=='f' and str(entry['cached']).startswith('#'):issues.append({'sheet':s.title,'cell':c.coordinate,'kind':'cached_error','detail':entry['cached']})
 inventory.append({'index':len(inventory),'name':s.title,'rows':s.max_row,'columns':s.max_column,'merged':[str(r) for r in s.merged_cells.ranges],'cells':cells,'role':'auxiliary' if s.title in ['Zalihe na parceli','Priručnik'] else 'production'})
 if s.title.startswith('Stočarstvo'):
  for namecol,sharecol,qtycol,vetcol in [('A','B','C','D'),('G','H','I','J')]:
   for r in [2,21,38,52,68,84]:
    name=s[f'{namecol}{r}'].value; head=f'{sharecol}{r}'; items=[]; groups=[]
    end=next((i for i in range(r+2,min(r+18,95)) if isinstance(s[f'{sharecol}{i}'].value,str) and s[f'{sharecol}{i}'].value.startswith('=SUM')),94)
    # Each SUM row defines its own independent feed block.
    for i in range(r+2,min(r+19,95)):
     f=s[f'{sharecol}{i}'].value
     if isinstance(f,str) and f.startswith('=SUM'):
      m=re.search(r'\(([A-Z]+)(\d+):[A-Z]+(\d+)\)',f)
      if m:
       start,stop=int(m[2]),int(m[3]); group=f'{sharecol}{i}'; groups.append({'id':group,'source':f,'shares':[f'{sharecol}{j}' for j in range(start,stop+1)]})
       for j in range(start,stop+1):
        label=s[f'{namecol}{j}'].value
        if not label:label=f'Neimenovano hranivo {namecol}{j}'
        share=number(s.title,f'{sharecol}{j}'); q=number(s.title,f'{qtycol}{j}',{head:1})
        items.append({'id':f'{qtycol}{j}','name':label,'quantity':q,'price':None,'unit':'kg','group':group,'share':share,'formula':s[f'{qtycol}{j}'].value,'quantityCell':f'{qtycol}{j}','priceCell':None,'status':'review','zeroReason':None})
    items.append({'id':f'{vetcol}{r}','name':'Veterinarski troškovi','quantity':'1','price':number(s.title,f'{vetcol}{r}'),'unit':'grlo','quantityCell':head,'priceCell':f'{vetcol}{r}','status':'review','zeroReason':None})
    cycle=name in ['Tovna junad','Prasad','Tovne svinje','Tovni pilići','Jagnjad','Jarad']
    weightrow={21:20,38:37,84:83}.get(r)
    weights=None
    if cycle and name in ['Tovna junad','Prasad','Tovne svinje','Tovni pilići']:
     formula=s[f'{"E" if namecol=="A" else "K"}{weightrow}'].value
     weights={'formula':formula,'buy':s[f'{sharecol}{weightrow}'].value if name!='Tovni pilići' else 0,'sell':s[f'{vetcol}{weightrow}'].value,'coefficient':str(formula).split('*')[-1] if formula else None}
    cards.append({'id':f's{w.sheetnames.index(s.title)}-{namecol}{r}','sheet':s.title,'name':name,'range':f'{namecol}{r}:{qtycol}{r+17}','basis':'head_cycle' if cycle else 'head_year','reference':'1','intensity':'niži' if 'niži' in s.title else 'viši','items':items,'groups':groups,'weights':weights,'notes':[],'issues':['Potvrditi vremensku osnovu i cene hraniva.'],'status':'archive'})
  continue
 if s.title in ['Zalihe na parceli','Priručnik']:continue
 headers=[]
 for row in s:
  r=row[0].row
  if s[f'B{r}'].value and (str(s[f'D{r}'].value).startswith('Obradivo') or (s.title=='Plastenici' and s[f'F{r}'].value=='Broj plastenika') or (s.title=='Ostalo' and s[f'B{r}'].value=='ŠAMPINJONI')):headers.append((r,s[f'B{r}'].value))
 if s.title=='Plastenici':headers += [(196,s['A196'].value),(219,s['A219'].value)]
 for n,(start,name) in enumerate(headers):
  stop=headers[n+1][0]-1 if n+1<len(headers) else s.max_row
  seed=s.title=='Plastenici' and start>=196; greenhouse=s.title=='Plastenici' and not seed
  qcol,pcol,totalcol=('C','D','E') if seed else ('D','E','F') if greenhouse else ('E','D','F')
  items=[]; notes=[]; localissues=[]; total=None
  for r in range(start+1,stop+1):
   label=s[f'B{r}'].value
   if 'UKUPN' in str(s[f'A{r}'].value).upper() or 'UKUPNO'==str(label):total=f'{totalcol}{r}'; break
  totalrow=int(re.search(r'\d+',total)[0]) if total else stop
  for r in range(start+1,totalrow):
   label=s[f'B{r}'].value; f=s[f'{totalcol}{r}'].value
   if not label or not (isinstance(f,str) and f.startswith('=')):continue
   qty=number(s.title,f'{qcol}{r}'); price=number(s.title,f'{pcol}{r}')
   flags=[]
   if qty is not None and Decimal(qty)>0 and price is None:flags.append('missing_price')
   if price=='0' and qty is not None and Decimal(qty)>0:flags.append('unconfirmed_zero')
   if ' ili' in str(label).lower() or (r>start+1 and ' ili' in str(s[f'B{r-1}'].value).lower()):flags.append('alternative')
   if re.search(r'\(\d/\d\)',str(label)):flags.append('allocation_already_in_source_quantity')
   if re.search(r'[IVX]+-[IVX]+',str(s[f'C{r}'].value)):flags.append('month_range')
   item={'id':f'{totalcol}{r}','name':label,'quantity':qty,'price':price,'unit':'kg/l — potvrditi' if not s[f'A{r}'].value else 'l goriva/usluga — potvrditi','month':s[f'C{r}'].value if not seed else None,'quantityCell':f'{qcol}{r}','priceCell':f'{pcol}{r}','formula':f,'originalQuantity':s[f'{qcol}{r}'].value,'originalPrice':s[f'{pcol}{r}'].value,'cachedCost':cached[s.title][f'{totalcol}{r}'].value,'flags':flags,'alternativeGroup':f'B{r}' if ' ili' in str(label).lower() else f'B{r-1}' if r>start+1 and ' ili' in str(s[f'B{r-1}'].value).lower() else None,'status':'review','zeroReason':None}
   items.append(item)
   for flag in flags:issues.append({'sheet':s.title,'cell':f'{pcol}{r}' if 'price' in flag or 'zero' in flag else f'B{r}','kind':flag,'detail':str(label)})
  for r in range(totalrow+1,stop+1):
   for c in s[r]:
    if isinstance(c.value,str) and len(c.value)>150 and not c.value.startswith('='):notes.append({'cell':c.coordinate,'text':c.value})
  if name=='ŠAMPINJONI':localissues.append('Neusaglašena jedinica: košnice / vreća. Admin mora potvrditi vreću.')
  if 'MALINA' in name:localissues.append('Zajednički normativ; obavezan izbor proizvoda malina ili kupina.')
  cards.append({'id':f's{w.sheetnames.index(s.title)}-B{start}','sheet':s.title,'name':name,'range':f'A{start}:I{stop}','basis':'seedlings' if seed else 'm2' if greenhouse else 'bag' if name=='ŠAMPINJONI' else 'ha','reference':'80000' if start==196 and seed else '5000' if seed else '100' if greenhouse else '1','items':items,'groups':[],'notes':notes,'issues':localissues,'totalCell':total,'cachedTotal':cached[s.title][total].value if total else None,'recalculatedHistorical':number(s.title,total) if total else None,'status':'archive'})
fixtures=[]
for sh,cell,expected in [('Usevi','F27','355.5774'),('Usevi','F62','406.3616'),('Plastenici','F24','162.25996'),('Plastenici','H24','0.4056499'),('Ostalo','F34','2.77693')]:
 actual=calc(sh,cell); assert abs(actual-Decimal(expected))<=Decimal('.000001'),(sh,cell,actual)
 fixtures.append({'sheet':sh,'cell':cell,'expected':expected,'calculated':str(actual),'passed':True})
result={'name':'Izvor 2017 — arhivski','sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'extractedAt':now,'sheets':inventory,'cards':cards,'issues':issues,'fixtures':fixtures}
(ROOT/'source/inventory.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
lines=['# Izvorni audit — Izvor 2017 — arhivski','',f'SHA-256 originala: `{result["sha256"]}`. Original je kopiran bajt-po-bajt; izvorni prilog nije izmenjen.','', '| Izvorni sheet | Uloga | Karte | Formule |','|---|---|---:|---:|']
for s in inventory:lines.append(f'| {s["name"]} | {s["role"]} | {sum(c["sheet"]==s["name"] for c in cards)} | {sum(bool(c["formula"]) for c in s["cells"])} |')
lines += ['',f'Ukupno {len(cards)} predložaka; {sum(i["kind"]=="external" for i in issues)} spoljnih formula. Sve ćelije, izvorne formule, keširane vrednosti, komentari i zavisnosti su u source/inventory.json. Uvoz ne izvršava formule.','', '## Regresija istorijskih rezultata','']+[f'- {f["sheet"]}!{f["cell"]}: {f["calculated"]} EUR — prolazi toleranciju 0,000001.' for f in fixtures]
lines+=['','## Pravila i otvorena pitanja','', '- Sve cene 2017 su arhiva. Nema aktivnog cenovnika. Administrator klonira, potvrđuje jedinice, grupe, vremensku osnovu i alternative i odobrava cene.','- Prazna cena pri pozitivnoj količini blokira aktiviranje. Izričita nula zahteva razlog. Neprimenljiva količina nula ne zahteva cenu.','- Alternative „ili“ se ne biraju automatski. Administrator mora eksplicitno odabrati uključenu stavku ili potvrditi isključenje.','- Šampinjoni: košnice u E23:F23, vreće u A37 i Priručniku. Vreća je predlog; do potvrde aktivacija blokirana.','- Normativne raspodele 1/6 i 1/2 su već u izvornim količinama: ne primenjivati ih drugi put i ne stvarati plaćanja.','- Stočarski blokovi hraniva čuvaju zasebne SUM opsege. Cene hraniva nisu cene gotovog nacionalnog tržišta. Junice i ostale vremenske osnove zahtevaju admin potvrdu.','- Spoljne reference se mapiraju na stvarnu površinu i prinos gazdinstva. Nedostajući prinos je nepoznat podatak. Hektarski deficit bez prinosa prikazuje „Nedostaje prinos“.','- Zalihe na parceli ostaju pomoćni kumulativni modul; Priručnik je pregled postojećih normativa, bez duplog uvoza.','- Istorijske napomene i preparati su samo za admin pregled, nisu savremeni agronomski saveti.','', '## Obrada troškova','', 'Količina × odobrena cena; puna decimalna preciznost do prikaza. ha → ostvarena ha, plastenik → m²/100, rasad → biljke/80000 ili /5000. Godišnji broj grla koristi prosečan broj i stvarni udeo godine; ciklus koristi broj završenih grla, bez godišnjeg skaliranja. Pokriveni su samo materijal, gorivo i odobrene izvorne operacije; ne puna dobit, rad, amortizacija ili porezi.']
(ROOT/'docs/IZVORNI-AUDIT.md').write_text('\n'.join(lines))
print(json.dumps({'sheets':len(inventory),'cards':len(cards),'issues':len(issues),'external':sum(i['kind']=='external' for i in issues),'fixtures':fixtures},ensure_ascii=False))
